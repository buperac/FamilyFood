import base64
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2C2620")
WRAP = Alignment(vertical="top", wrap_text=True)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD
        cell.fill = FILL
    ws.freeze_panes = "A2"

wb = Workbook()

# ---------- Meal Log ----------
ml = wb.active
ml.title = "Meal Log"
ml_head = ["Date cooked", "Week", "Day", "Meal", "Cuisine", "Cal/adult",
           "Est. cost", "Actual cost", "Rating (1-5)", "Favourite", "Notes"]
ml.append(ml_head)
meals = [
 (1,"Mon","Sheet-Pan Chicken Fajitas","Tex-Mex",580,19.54),
 (1,"Tue","Ground Beef Bolognese + Caesar","Italian",830,9.22),
 (1,"Wed","Beef Kofta in Tikka Masala","Indian",620,19.68),
 (1,"Thu","Ground Beef Tacos","Tex-Mex",600,9.12),
 (1,"Fri","Build-Your-Own Pizza Night","Italian",820,13.93),
 (1,"Sat","Korean-Style Ground Beef Bowls","Asian",590,16.86),
 (1,"Sun","Roast Chicken + Potatoes","Classic",640,26.95),
 (2,"Mon","Cottage Pie","British",640,27.89),
 (2,"Tue","Chicken Stir-Fry + Rice","Asian",560,21.71),
 (2,"Wed","Sloppy Joes + Oven Fries","American",600,7.82),
 (2,"Thu","Honey-Garlic Chicken Thighs","Asian",580,18.37),
 (2,"Fri","Calzone Night","Italian",850,13.76),
 (2,"Sat","Smashburgers + Oven Fries","American",720,6.46),
 (2,"Sun","Pork Tenderloin Roast","Classic",600,25.26),
 (3,"Mon","Italian Wedding Soup","Italian",520,24.75),
 (3,"Tue","Chicken Quesadillas","Tex-Mex",760,24.16),
 (3,"Wed","Chicken Souvlaki Bowls","Classic",580,37.14),
 (3,"Thu","BBQ Chicken Drumsticks","BBQ",620,26.31),
 (3,"Fri","Pizza Night","Italian",820,13.93),
 (3,"Sat","Beef Tacos (Elevated)","Tex-Mex",580,9.75),
 (3,"Sun","Roast Chicken (Encore)","Classic",640,27.55),
]
for w, day, name, cuisine, cal, cost in meals:
    ml.append(["", w, day, name, cuisine, cal, cost, "", "", "", ""])
for r in range(2, ml.max_row + 1):
    ml.cell(row=r, column=7).number_format = '"$"#,##0.00'
    ml.cell(row=r, column=8).number_format = '"$"#,##0.00'
widths = [12, 6, 6, 34, 12, 9, 10, 11, 11, 10, 40]
for i, wdt in enumerate(widths, 1):
    ml.column_dimensions[chr(64 + i)].width = wdt
style_header(ml, len(ml_head))

# ---------- Ingredient Prices ----------
ip = wb.create_sheet("Ingredient Prices")
ip_head = ["Ingredient", "Category", "Est. price", "Actual / last paid",
           "Store / brand", "Size", "Per-unit", "Date updated", "Notes"]
ip.append(ip_head)
# (name, category, est_price, sub_note, stash, have)
ings = [
 ("Ground beef (~12 lb, 8 meals)","Protein",0,"from your freezer stash",True,False),
 ("Chicken breast x2 packs","Protein",48.38,"Maple Leaf value pack",False,False),
 ("Boneless chicken thighs","Protein",14.21,"Maple Leaf - not halal",False,False),
 ("Chicken drumsticks","Protein",14.81,"Maple Leaf - not halal",False,False),
 ("Whole chicken x2","Protein",38.56,"two Sunday roasts",False,False),
 ("Pork tenderloin x2","Protein",20.76,"Maple Leaf",False,False),
 ("Eggs x2 (60 ct)","Dairy & Eggs",20.54,"Great Value",False,False),
 ("Milk 2% 4L x2","Dairy & Eggs",13,"Dairyland",False,False),
 ("Cheddar block x2","Dairy & Eggs",10.96,"Great Value",False,False),
 ("Mozzarella block x2","Dairy & Eggs",10.96,"Great Value",False,False),
 ("Greek yogurt (plain)","Dairy & Eggs",6,"",False,False),
 ("Grated parmesan","Dairy & Eggs",6,"",False,False),
 ("Sour cream","Dairy & Eggs",3.5,"",False,False),
 ("Feta","Dairy & Eggs",6,"",False,False),
 ("Cottage cheese","Dairy & Eggs",5,"",False,False),
 ("Salted butter","Dairy & Eggs",5.5,"finishing / veg / toast",False,False),
 ("Unsalted butter","Dairy & Eggs",5.5,"sauces & baking",False,False),
 ("Frozen mixed berries","Frozen",13,"",False,False),
 ("Frozen peas","Frozen",6,"",False,False),
 ("Frozen broccoli","Frozen",6,"",False,False),
 ("Frozen corn","Frozen",6,"",False,False),
 ("Bell peppers x2 (3-packs)","Produce",13,"",False,False),
 ("Yellow onions (3 lb)","Produce",3,"",False,False),
 ("Yellow potatoes x2 (5 lb)","Produce",8,"",False,False),
 ("Carrots (3 lb)","Produce",3,"",False,False),
 ("Garlic x3 bulbs","Produce",3,"",False,False),
 ("Romaine hearts (3-pack)","Produce",5,"",False,False),
 ("English cucumbers x3","Produce",4.5,"",False,False),
 ("Cilantro","Produce",1.5,"",False,False),
 ("Lemons","Produce",4,"",False,False),
 ("Limes","Produce",2,"",False,False),
 ("Fresh ginger","Produce",2,"",False,False),
 ("Cherry tomatoes","Produce",4,"",False,False),
 ("Cooked beets","Produce",7,"lunch salads",False,False),
 ("Avocados","Produce",3,"",False,False),
 ("Coleslaw mix","Produce",3,"",False,False),
 ("Green beans","Produce",4,"",False,False),
 ("Fresh herbs (thyme/mint/basil)","Produce",6,"",False,False),
 ("Baby spinach","Produce",4,"",False,False),
 ("Jasmine rice 8 kg","Pantry & Dry",19.97,"lasts months",False,False),
 ("Rolled oats","Pantry & Dry",5,"",False,False),
 ("Granola","Pantry & Dry",5,"",False,False),
 ("Taco seasoning","Pantry & Dry",2,"",False,False),
 ("Marinara sauce","Pantry & Dry",4,"",False,False),
 ("Pizza sauce","Pantry & Dry",3,"",False,False),
 ("Tikka masala sauce","Pantry & Dry",4,"Patak's, mild",False,False),
 ("Canned beans x3","Pantry & Dry",5,"",False,False),
 ("Broth x2","Pantry & Dry",6,"",False,False),
 ("BBQ sauce","Pantry & Dry",4,"",False,False),
 ("Kalamata olives","Pantry & Dry",5,"lunch salads",False,False),
 ("Farro / quinoa","Pantry & Dry",5,"lunch salads",False,False),
 ("Walnuts","Pantry & Dry",6,"lunch salads",False,False),
 ("Chickpeas (canned)","Pantry & Dry",2,"lunch salads",False,False),
 ("Flour tortillas x2","Bakery",5,"",False,False),
 ("Pizza dough x3","Bakery",12,"",False,False),
 ("Naan","Bakery",4,"",False,False),
 ("Pepperoni","Bakery",5,"",False,False),
 ("Burger buns x2","Bakery",6,"",False,False),
 ("Spaghetti","Bakery",3,"",False,False),
 ("Orzo / acini","Bakery",3,"",False,False),
 ("Crusty bread","Bakery",4,"",False,False),
 ("Breakfast sausage","Bakery",7,"",False,False),
 ("Olive oil","Pantry Staples",6,"",False,True),
 ("Soy sauce","Pantry Staples",3,"",False,True),
 ("Sesame oil","Pantry Staples",5,"",False,True),
 ("Hoisin","Pantry Staples",4,"",False,True),
 ("Honey","Pantry Staples",5,"",False,True),
 ("Brown sugar","Pantry Staples",3,"",False,True),
 ("Worcestershire","Pantry Staples",3,"",False,True),
 ("Dijon mustard","Pantry Staples",3,"",False,True),
 ("Breadcrumbs","Pantry Staples",3,"",False,True),
 ("Garam masala","Pantry Staples",4,"",False,True),
 ("Spice rack (cumin, paprika, oregano)","Pantry Staples",0,"",True,True),
]
for name, cat, price, sub, stash, have in ings:
    note = sub
    flag = "freezer stash ($0)" if stash else ("owned pantry staple" if have else "")
    if flag:
        note = (sub + " - " if sub else "") + flag
    ip.append([name, cat, price, "", "", "", "", "", note])
for r in range(2, ip.max_row + 1):
    ip.cell(row=r, column=3).number_format = '"$"#,##0.00'
    ip.cell(row=r, column=4).number_format = '"$"#,##0.00'
ipw = [32, 14, 10, 16, 16, 12, 10, 12, 28]
for i, wdt in enumerate(ipw, 1):
    ip.column_dimensions[chr(64 + i)].width = wdt
style_header(ip, len(ip_head))

# ---------- Deals ----------
dl = wb.create_sheet("Deals")
dl_head = ["Date", "Item", "Store", "Sale per-unit", "Baseline per-unit",
           "Delta %", "Shelf-stable?", "Action / qty grabbed"]
dl.append(dl_head)
dlw = [12, 30, 12, 14, 16, 9, 14, 26]
for i, wdt in enumerate(dlw, 1):
    dl.column_dimensions[chr(64 + i)].width = wdt
style_header(dl, len(dl_head))

path = r"C:\Users\kyle\Recipes\Gronning-Kitchen-Tracker.xlsx"
wb.save(path)
with open(path, "rb") as f:
    data = f.read()
with open(r"C:\Users\kyle\Recipes\_tracker.b64", "w") as f:
    f.write(base64.b64encode(data).decode())
print("OK", len(data), "bytes;", len(meals), "meals;", len(ings), "ingredients")
